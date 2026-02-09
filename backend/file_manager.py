import os
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor
import logging

logger = logging.getLogger(__name__)

def save_meeting_log(save_dir, session_title, transcripts, start_time=None, summary=None, advices=None):
    """会議ログをMarkdownファイルとして保存
    
    Args:
        save_dir: 保存先ディレクトリ
        session_title: 会議タイトル
        transcripts: 文字起こしデータ
        start_time: 会議開始日時（ISO形式文字列）
        summary: サマリーテキスト
        advices: アドバイスリスト
    """
    try:
        # ディレクトリ作成
        save_path = Path(save_dir)
        save_path.mkdir(parents=True, exist_ok=True)
        
        # 日時情報を取得
        if start_time:
            try:
                dt = datetime.fromisoformat(start_time.replace('Z', '+00:00'))
            except:
                dt = datetime.now()
        else:
            dt = datetime.now()
        
        # ファイル名生成 (タイムスタンプ + タイトル)
        timestamp = dt.strftime("%Y%m%d_%H%M%S")
        safe_title = "".join([c for c in session_title if c.isalnum() or c in (' ', '-', '_')]).strip()
        filename_base = f"{timestamp}_{safe_title}"
        md_file_path = save_path / f"{filename_base}.md"
        
        with open(md_file_path, 'w', encoding='utf-8') as f:
            f.write(f"# {session_title}\n\n")
            f.write(f"**Date:** {dt.strftime('%Y-%m-%d %H:%M')}\n\n")
            
            if summary:
                f.write("## 📝 Summary\n\n")
                f.write(summary + "\n\n")
                f.write("---\n\n")
            
            f.write("## 💬 Transcript\n\n")
            for t in transcripts:
                # タイムスタンプの整形 (例: 2024-02-06 10:00:00 -> 10:00:00)
                try:
                    time_str = datetime.fromisoformat(t['timestamp']).strftime('%H:%M:%S')
                except:
                    time_str = t['timestamp']
                
                icon = "🔊" if t['source'] == 'speaker' else "🎤"
                role = "Speaker" if t['source'] == 'speaker' else "Mic"
                
                f.write(f"**{icon} {role}** ({time_str}):\n")
                f.write(f"{t['text']}\n\n")
                
            if advices:
                f.write("\n---\n\n")
                f.write("## 🤖 AI Advice Log\n\n")
                for a in advices:
                    f.write(f"- **Trigger:** {a['trigger_condition']}\n")
                    f.write(f"  - {a['advice_text']}\n")

        logger.info(f"Saved Markdown log to: {md_file_path}")
        return str(md_file_path)
        
    except Exception as e:
        logger.error(f"Failed to save markdown log: {e}")
        return None

def generate_docx(session_title, transcripts, start_time=None, summary=None, advices=None):
    """会議ログのDocxオブジェクトを生成
    
    Args:
        session_title: 会議タイトル
        transcripts: 文字起こしデータ
        start_time: 会議開始日時（ISO形式文字列）
        summary: サマリーテキスト
        advices: アドバイスリスト
    """
    doc = Document()
    
    # 日時情報を取得
    if start_time:
        try:
            dt = datetime.fromisoformat(start_time.replace('Z', '+00:00'))
        except:
            dt = datetime.now()
    else:
        dt = datetime.now()
    
    # タイトル
    doc.add_heading(session_title, 0)
    doc.add_paragraph(f"Date: {dt.strftime('%Y-%m-%d %H:%M')}")
    
    # サマリー
    if summary:
        doc.add_heading('Summary', level=1)
        doc.add_paragraph(summary)
    
    # トランスクリプト
    doc.add_heading('Transcript', level=1)
    
    for t in transcripts:
        try:
            time_str = datetime.fromisoformat(t['timestamp']).strftime('%H:%M:%S')
        except:
            time_str = t['timestamp']
            
        p = doc.add_paragraph()
        run = p.add_run(f"[{time_str}] {t['source'].upper()}: ")
        run.bold = True
        if t['source'] == 'speaker':
            run.font.color.rgb = RGBColor(0, 100, 0) # Dark Green
        else:
            run.font.color.rgb = RGBColor(0, 0, 139) # Dark Blue
            
        p.add_run(t['text'])
        
    # アドバイス
    if advices:
        doc.add_heading('AI Advice Log', level=1)
        for a in advices:
            p = doc.add_paragraph(style='List Bullet')
            trigger = p.add_run(f"Trigger: {a['trigger_condition']}")
            trigger.bold = True
            doc.add_paragraph(a['advice_text'], style='List Continue')
            
    return doc
